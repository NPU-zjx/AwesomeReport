"""Rebuild the sanitized standard template from its checked-in typography profile."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from format_report import ROOT, DEFAULT_PROFILE, apply_profile, field, borders


def create(output, force=False):
    output=Path(output)
    if output.exists() and not force: raise FileExistsError('Template exists; use --force to rebuild deliberately')
    doc=Document()
    apply_profile(doc,json.loads(DEFAULT_PROFILE.read_text(encoding='utf-8')))
    for key in ('author','last_modified_by','subject','comments','keywords','category'):
        setattr(doc.core_properties,key,'')
    doc.core_properties.title='中文报告标准模板'
    doc.add_paragraph('报告标题',style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('1 一级标题示例',style='Heading 1')
    doc.add_paragraph('此处填写报告正文。正文中文使用宋体，英文与数字使用Times New Roman。章节名称与内容由当前任务决定，本模板不包含实验结论或作者身份。')
    doc.add_paragraph('1.1 二级标题示例',style='Heading 2')
    doc.add_paragraph('本文件提供可编辑的段落、表格和代码样式。正文首行缩进两字，使用1.25倍行距。请按指定任务书或模板覆盖相应格式。')
    p=doc.add_paragraph('表1 表格样式示例',style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=2,cols=2); borders(table)
    for row,values in zip(table.rows,[['项目','说明'],['示例单元格','此处填写真实数据或说明']]):
        for cell,value in zip(row.cells,values):
            cell.paragraphs[0].style='Table Text'; cell.paragraphs[0].add_run(value)
    doc.add_paragraph('代码样式示例（非实验结果）',style='Heading 2')
    table=doc.add_table(rows=1,cols=1); borders(table)
    p=table.cell(0,0).paragraphs[0]; p.style='Code'
    p.add_run('def example(value):\n    return value')
    doc.add_paragraph('图片样式示例', style='Heading 2')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/'examples/workflow.png'),width=Cm(14))
    p=doc.add_paragraph('图1 通用处理流程（仅示意）',style='Caption')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.sections[0].footer.paragraphs[0]; p.style='Footer'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(p,'PAGE')
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',type=Path,default=ROOT/'assets/standard-report.docx')
    p.add_argument('--force',action='store_true')
    a=p.parse_args(); create(a.output,a.force)
