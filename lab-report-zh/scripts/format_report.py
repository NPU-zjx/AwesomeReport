"""Build an editable DOCX from explicit JSON blocks and a retained Word template.

Requires python-docx. No Word, network or PDF converter is needed to build.
This command creates a new report; it never reformats an existing report in place.
"""
import argparse
import json
from pathlib import Path
import tempfile
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / 'assets/standard-report.docx'
DEFAULT_PROFILE = ROOT / 'assets/typography.json'


def font_properties(element, chinese, latin, size=None, bold=None):
    pr = element.get_or_add_rPr()
    rf = pr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); pr.insert(0, rf)
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    for key in ('ascii', 'hAnsi', 'cs'): rf.set(qn('w:' + key), latin)
    rf.set(qn('w:eastAsia'), chinese)
    if size is not None:
        for name in ('sz', 'szCs'):
            child = pr.find(qn('w:' + name))
            if child is None: child = OxmlElement('w:' + name); pr.append(child)
            child.set(qn('w:val'), str(round(size * 2)))
    if bold is not None:
        for name in ('b', 'bCs'):
            child = pr.find(qn('w:' + name))
            if child is None: child = OxmlElement('w:' + name); pr.append(child)
            child.set(qn('w:val'), '1' if bold else '0')


def apply_profile(doc, profile):
    """Apply a deliberately selected profile, not an implicit external-template override."""
    from docx.enum.style import WD_STYLE_TYPE
    for name, values in profile['styles'].items():
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st = doc.styles[name]
        st.font.color.rgb = RGBColor(0, 0, 0)
        if name != 'Normal': st.base_style = doc.styles['Normal']
        font_properties(st.element, values.get('chinese', profile['chinese']),
                        profile['latin'], values['size'], values.get('bold', False))
        pf = st.paragraph_format
        pf.page_break_before = False
        pf.first_line_indent = Pt(values.get('indent', 0))
        pf.line_spacing = values.get('line_spacing', 1.25)
        pf.space_before = Pt(values.get('before', 0))
        pf.space_after = Pt(values.get('after', 0))
        pf.keep_with_next = name.startswith('Heading')
        pf.keep_together = False
        pf.widow_control = True
        # Remove decorative borders inherited from the runtime's default template.
        ppr = st.element.get_or_add_pPr()
        for node in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(node)
        for name_ in ('autoSpaceDE', 'autoSpaceDN'):
            ppr = st.element.get_or_add_pPr()
            child = ppr.find(qn('w:' + name_))
            if child is None: child = OxmlElement('w:' + name_); ppr.append(child)
            child.set(qn('w:val'), '0')
    page = profile['page']
    for sec in doc.sections:
        sec.page_width = Cm(page['width']); sec.page_height = Cm(page['height'])
        for key in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin',
                    'header_distance', 'footer_distance'):
            setattr(sec, key, Cm(page[key]))


def field(paragraph, instruction, placeholder=''):
    r = paragraph.add_run()._r
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); r.append(begin)
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve')
    code.text = ' ' + instruction + ' '; r.append(code)
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); r.append(sep)
    if placeholder: paragraph.add_run(placeholder)
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); paragraph.add_run()._r.append(end)


def borders(table):
    pr = table._tbl.tblPr
    node = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + key), value)
        node.append(edge)
    pr.append(node)


def add_caption(doc, text, before=False):
    p = doc.add_paragraph(text, style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = before
    return p


def validate(data, base):
    if not isinstance(data, dict) or not isinstance(data.get('blocks'), list):
        raise ValueError('Input must be an object with a blocks list')
    for key in ('title', 'author', 'subject', 'date'):
        if key in data and not isinstance(data[key], str): raise ValueError(key + ' must be text')
    for key in ('toc', 'cover'):
        if key in data and not isinstance(data[key], bool): raise ValueError(key + ' must be boolean')
    if not isinstance(data.get('metadata', {}), dict): raise ValueError('metadata must be an object')
    if not all(isinstance(k, str) and isinstance(v, str) for k, v in data.get('metadata', {}).items()):
        raise ValueError('metadata keys and values must be text')
    for i, b in enumerate(data['blocks']):
        if not isinstance(b, dict): raise ValueError(f'Block {i} must be an object')
        kind = b.get('type')
        if kind not in ('paragraph', 'heading', 'code', 'table', 'image', 'page_break'):
            raise ValueError(f'Unknown block type at {i}: {kind}')
        if kind in ('paragraph', 'heading', 'code') and not isinstance(b.get('text'), str):
            raise ValueError(f'Block {i} needs text')
        if kind == 'heading' and (type(b.get('level', 1)) is not int or not 1 <= b.get('level', 1) <= 6):
            raise ValueError('Heading level must be 1..6')
        if 'caption' in b and not isinstance(b['caption'], str): raise ValueError('caption must be text')
        if kind == 'table':
            rows = b.get('rows')
            if not isinstance(rows, list) or not rows or not isinstance(rows[0], list) or not rows[0]:
                raise ValueError('Table rows must be a nonempty rectangular list')
            if any(not isinstance(r, list) or len(r) != len(rows[0]) or
                   not all(isinstance(c, str) for c in r) for r in rows):
                raise ValueError('Table cells must be text and rows have equal length')
        if kind == 'image':
            if not isinstance(b.get('path'), str) or not (base / b['path']).is_file():
                raise ValueError(f'Image missing in block {i}')
            if 'width_cm' in b and (not isinstance(b['width_cm'], (int, float)) or b['width_cm'] <= 0):
                raise ValueError('Image width must be positive')


def build(input_path, output_path, template=DEFAULT_TEMPLATE, style_config=None, force=False):
    input_path, output_path, template = map(lambda p: Path(p).resolve(), (input_path, output_path, template))
    if output_path.suffix.lower() != '.docx': raise ValueError('Output must end with .docx')
    data = json.loads(input_path.read_text(encoding='utf-8-sig'))
    validate(data, input_path.parent)
    protected = {input_path, template}
    protected.update((input_path.parent/b['path']).resolve() for b in data['blocks'] if b['type']=='image')
    if style_config: protected.add(Path(style_config).resolve())
    if output_path in protected or any(output_path.exists() and output_path.samefile(p) for p in protected):
        raise ValueError('Output must not overwrite an input, template or image')
    if output_path.exists() and not force: raise FileExistsError('Output exists; use a new name or --force')
    doc = Document(template)
    # A content template is a style/page reference. Its old body is not report evidence.
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'): body.remove(child)
    if style_config:
        apply_profile(doc, json.loads(Path(style_config).read_text(encoding='utf-8-sig')))
    for required in ('Normal', 'Title', 'Caption'):
        if required not in doc.styles: raise ValueError('Template missing style: '+required)
    for prop in ('author', 'last_modified_by', 'subject', 'title', 'comments', 'keywords', 'category'):
        setattr(doc.core_properties, prop, '')
    doc.core_properties.title = data.get('title', '')
    doc.core_properties.author = data.get('author', '')
    doc.core_properties.subject = data.get('subject', '')
    if data.get('title'):
        doc.add_paragraph(data['title'], style='Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, value in data.get('metadata', {}).items():
        p = doc.add_paragraph(f'{label}：{value}')
        p.paragraph_format.first_line_indent = Pt(0)
    if data.get('date'): doc.add_paragraph('日期：'+data['date'])
    if data.get('cover'): doc.add_page_break()
    if data.get('toc'):
        p = doc.add_paragraph('目录', style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', '请在Word中更新目录')
        doc.add_page_break()
    width = doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin
    width_cm = width / 360000
    for b in data['blocks']:
        kind = b['type']
        if kind == 'paragraph': doc.add_paragraph(b['text'])
        elif kind == 'heading': doc.add_paragraph(b['text'], style=f'Heading {b.get("level", 1)}')
        elif kind == 'page_break': doc.add_page_break()
        elif kind == 'code':
            if 'Code' not in doc.styles: raise ValueError('Template needs a Code style for code blocks')
            if b.get('caption'): add_caption(doc, b['caption'], True)
            table = doc.add_table(rows=1, cols=1); borders(table)
            table.columns[0].width = width
            lines = b['text'].split('\n')
            for i, line in enumerate(lines):
                p = table.cell(0, 0).paragraphs[0] if i == 0 else table.cell(0, 0).add_paragraph()
                p.style = 'Code'; p.add_run(line)
                p.paragraph_format.keep_with_next = len(lines) <= 12 and i < len(lines)-1
                p.paragraph_format.keep_together = True
        elif kind == 'table':
            if 'Table Text' not in doc.styles: raise ValueError('Template needs a Table Text style')
            if b.get('caption'): add_caption(doc, b['caption'], True)
            table = doc.add_table(rows=len(b['rows']), cols=len(b['rows'][0])); borders(table)
            for col in table.columns: col.width = int(width/len(table.columns))
            for i, row in enumerate(b['rows']):
                if i == 0:
                    header = OxmlElement('w:tblHeader'); table.rows[i]._tr.get_or_add_trPr().append(header)
                for j, text in enumerate(row):
                    p = table.cell(i, j).paragraphs[0]; p.style = 'Table Text'; r = p.add_run(text)
                    if i == 0: r.bold = True
        elif kind == 'image':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = bool(b.get('caption'))
            shape = p.add_run().add_picture(str(input_path.parent/b['path']), width=Cm(min(width_cm, b.get('width_cm',width_cm))))
            max_height = doc.sections[-1].page_height-doc.sections[-1].top_margin-doc.sections[-1].bottom_margin-Cm(2)
            if shape.height > max_height:
                shape.width = round(shape.width*max_height/shape.height); shape.height = max_height
            shape._inline.docPr.set('descr', b.get('caption', ''))
            if b.get('caption'): add_caption(doc, b['caption'])
    settings = doc.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None: update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'), 'true')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(suffix='.docx', dir=output_path.parent); os.close(fd)
    try:
        doc.save(temporary)
        os.replace(temporary, output_path)
    finally:
        if os.path.exists(temporary): os.unlink(temporary)
    return output_path


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--template', type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument('--style-config', type=Path, help='Explicitly override template styles with a full profile')
    parser.add_argument('--force', action='store_true')
    args = parser.parse_args()
    try:
        result = build(args.input,args.output,args.template,args.style_config,args.force)
    except (ValueError, OSError, KeyError) as exc:
        parser.exit(2, f'ERROR: {exc}\n')
    print(result)
    print('DOCX generated. Refresh fields and visually review before delivery.')
