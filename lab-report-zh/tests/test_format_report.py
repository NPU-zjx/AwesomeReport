import base64
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from format_report import build, DEFAULT_TEMPLATE, DEFAULT_PROFILE


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root/'内容.json'
        self.output = self.root/'结果.docx'

    def write(self, blocks, **extra):
        self.input.write_text(json.dumps({'blocks':blocks, **extra},ensure_ascii=False),encoding='utf-8')

    def test_content_styles_and_optional_toc(self):
        before=hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).digest()
        code='if value:\n    print("中文 123")\n\n    return value'
        self.write([{'type':'heading','level':1,'text':'12 自定义章节'},
                    {'type':'paragraph','text':'中英文 A B 与必要空格'},
                    {'type':'code','text':code},
                    {'type':'table','rows':[['列一','列二'],['甲','乙']]}],title='测试',toc=True)
        build(self.input,self.output)
        doc=Document(self.output)
        self.assertIn('12 自定义章节',[p.text for p in doc.paragraphs])
        self.assertEqual('\n'.join(p.text for p in doc.tables[0].cell(0,0).paragraphs),code)
        self.assertEqual(doc.styles['Normal'].font.size.pt,12)
        self.assertEqual(doc.styles['Code'].font.size.pt,9.5)
        self.assertEqual(doc.styles['Normal'].element.rPr.rFonts.get(qn('w:eastAsia')),'宋体')
        self.assertEqual(doc.styles['Normal'].element.rPr.rFonts.get(qn('w:ascii')),'Times New Roman')
        self.assertFalse(doc.styles['Heading 1'].paragraph_format.page_break_before)
        self.assertEqual(len(doc.element.xpath('.//w:br[@w:type="page"]')),1)
        self.assertTrue(any('TOC' in (n.text or '') for n in doc.element.xpath('.//w:instrText')))
        self.assertEqual(hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).digest(),before)

    def test_no_synthetic_identity_or_toc(self):
        self.write([{'type':'paragraph','text':'内容'}])
        build(self.input,self.output)
        d=Document(self.output)
        self.assertEqual([p.text for p in d.paragraphs],['内容'])
        self.assertEqual(d.core_properties.author,'')
        self.assertFalse(d.element.xpath('.//w:instrText'))

    def test_overwrite_and_input_protection(self):
        self.write([])
        self.output.write_bytes(b'keep')
        with self.assertRaises(FileExistsError): build(self.input,self.output)
        self.assertEqual(self.output.read_bytes(),b'keep')
        with self.assertRaises(ValueError): build(self.input,DEFAULT_TEMPLATE,force=True)
        build(self.input,self.output,force=True)
        Document(self.output)

    def test_invalid_data_does_not_replace_output(self):
        self.output.write_bytes(b'keep')
        cases=[{'type':'heading','level':0,'text':'bad'},
               {'type':'table','rows':[['a'],['b','c']]},
               {'type':'image','path':'missing.png'}, {'type':'unknown'}]
        for block in cases:
            self.write([block])
            with self.assertRaises(ValueError): build(self.input,self.output,force=True)
            self.assertEqual(self.output.read_bytes(),b'keep')

    def test_custom_template_preserved_unless_explicit_override(self):
        custom=self.root/'custom.docx'
        d=Document(DEFAULT_TEMPLATE); d.styles['Normal'].font.size=Pt(15); d.save(custom)
        self.write([{'type':'paragraph','text':'正文'}])
        build(self.input,self.output,custom)
        self.assertEqual(Document(self.output).styles['Normal'].font.size.pt,15)
        build(self.input,self.output,custom,DEFAULT_PROFILE,True)
        self.assertEqual(Document(self.output).styles['Normal'].font.size.pt,12)

    def test_image_bytes_and_caption_preserved(self):
        raw=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZlS8AAAAASUVORK5CYII=')
        (self.root/'image.png').write_bytes(raw)
        self.write([{'type':'image','path':'image.png','width_cm':2,'caption':'图1 示例'}])
        build(self.input,self.output)
        d=Document(self.output)
        self.assertEqual(len(d.inline_shapes),1)
        self.assertIn('图1 示例',[p.text for p in d.paragraphs])
        with ZipFile(self.output) as z:
            self.assertIn(raw,[z.read(n) for n in z.namelist() if n.startswith('word/media/')])


if __name__=='__main__': unittest.main()
